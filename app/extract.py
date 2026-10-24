"""Text extraction for uploaded DOCX / PDF / TXT. Runs locally; files never leave the server."""
import base64
import io


def extract_text(filename, b64data):
    raw = base64.b64decode(b64data)
    name = (filename or "").lower()
    if name.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(raw))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n\n".join(parts)
    if name.endswith(".pdf"):
        import pdfplumber
        with pdfplumber.open(io.BytesIO(raw)) as pdf:
            pages = [(p.extract_text() or "") for p in pdf.pages]
        return "\n\n".join(x.strip() for x in pages if x.strip())
    if name.endswith((".txt", ".md")):
        return raw.decode("utf-8", errors="replace")
    raise ValueError("Unsupported file type. Upload .docx, .pdf or .txt")
